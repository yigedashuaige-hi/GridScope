from __future__ import annotations

import argparse
import re
import shutil
from datetime import datetime, time
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import csr_matrix, lil_matrix


# ============================================================
# 1. 默认文件位置与模型参数
# ============================================================
DATA_CANDIDATES = [
    Path("附件1.xlsx"),
    Path("附件1_数据处理.xlsx"),
    Path("附件一_数据处理.xlsx"),
    Path("upload/附件1.xlsx"),
]

TEMPLATE_CANDIDATES = [
    Path("result1.xlsx"),
    Path("result1(1).xlsx"),
    Path("upload/result1.xlsx"),
]

DEFAULT_OUTPUT_DIR = Path("问题1输出")

N = 144                    # 24 h × 6 = 144 个10分钟时段
DELTA_T = 10.0 / 60.0      # h
ETA_C = 0.90               # 单次充电效率
ETA_D = 0.90               # 单次放电效率
E_MIN = 1200.0             # kWh
E_MAX = 10800.0            # kWh
P_MAX = 5000.0             # kW，按储能与母线接口侧理解
Q_MAX = P_MAX * DELTA_T    # 833.333333 kWh/10 min

E0_LIST = [1200, 3600, 6000, 8400, 10800]
BASE_E0 = 6000             # 正式 result1.xlsx 使用的基准初始电量

ZERO_TOL = 1e-7
CHECK_TOL = 1e-5


def find_existing_file(
    explicit_path: Path | None,
    candidates: list[Path],
    description: str,
) -> Path:
    """优先使用命令行指定文件，否则依次查找候选文件。"""
    if explicit_path is not None:
        if explicit_path.exists():
            return explicit_path
        raise FileNotFoundError(f"指定的{description}不存在：{explicit_path}")

    for path in candidates:
        if path.exists():
            return path

    names = "、".join(str(path) for path in candidates)
    raise FileNotFoundError(
        f"没有找到{description}。请确认脚本与文件位于同一文件夹，"
        f"或用命令行参数明确指定。\n程序尝试查找：{names}"
    )


# ============================================================
# 2. 时间解析：保留 day offset，不再对1440取模
# ============================================================
TIME_PATTERN = re.compile(
    r"^(?P<hour>\d{1,2}):(?P<minute>\d{1,2})(?:\+(?P<day>\d+))?$"
)


def excel_time_to_absolute_minutes(value) -> int:
    """
    将 Excel 时间转换为绝对分钟。

    示例：
        0:10    -> 10
        23:50   -> 1430
        0:00+1  -> 1440

    关键点：绝不使用 ``% 1440`` 抹去 ``+1`` 的日期偏移。
    """
    if value is None:
        raise ValueError("发现空时间单元格")

    if isinstance(value, datetime):
        return value.hour * 60 + value.minute

    if isinstance(value, time):
        return value.hour * 60 + value.minute

    if isinstance(value, (int, float, np.integer, np.floating)):
        numeric = float(value)
        if not np.isfinite(numeric):
            raise ValueError(f"时间不是有限数：{value}")
        # 时间型 Excel 数值通常是一天的小数；若显式大于等于1，则保留日偏移。
        return int(round(numeric * 1440.0))

    text = str(value).strip().replace(" ", "")
    match = TIME_PATTERN.fullmatch(text)
    if match is None:
        try:
            numeric = float(text)
        except ValueError as exc:
            raise ValueError(f"无法识别时间：{value}") from exc
        return int(round(numeric * 1440.0))

    hour = int(match.group("hour"))
    minute = int(match.group("minute"))
    day_offset = int(match.group("day") or 0)
    if not (0 <= hour <= 23 and 0 <= minute <= 59):
        raise ValueError(f"时间超出合法范围：{value}")
    return day_offset * 1440 + hour * 60 + minute


def unwrap_time_sequence(raw_minutes: list[int]) -> np.ndarray:
    """把跨午夜但未显式带 ``+1`` 的时间序列展开为严格递增绝对分钟。"""
    absolute: list[int] = []
    for raw in raw_minutes:
        minute = int(raw)
        while absolute and minute <= absolute[-1]:
            minute += 1440
        absolute.append(minute)

    result = np.asarray(absolute, dtype=int)
    if len(result) > 1 and np.any(np.diff(result) != 10):
        pairs = [
            f"{absolute_minute_text(result[i])}->{absolute_minute_text(result[i + 1])}"
            for i in range(len(result) - 1)
            if result[i + 1] - result[i] != 10
        ]
        raise ValueError(
            "时间序列必须按10分钟严格递增，异常相邻时刻：" + "、".join(pairs[:8])
        )
    return result


def parse_interval_start(interval_text) -> int:
    """读取模板时间段起点，如 ``10:00-10:10`` 或 ``0:00+1-0:10+1``。"""
    text = str(interval_text).strip()
    if "-" not in text:
        raise ValueError(f"无法识别时间段：{interval_text}")
    return excel_time_to_absolute_minutes(text.split("-", 1)[0].strip())


def absolute_minute_text(minute: int) -> str:
    day_offset, minute_of_day = divmod(int(minute), 1440)
    hour, minute_part = divmod(minute_of_day, 60)
    suffix = f"+{day_offset}" if day_offset else ""
    return f"{hour}:{minute_part:02d}{suffix}"


# ============================================================
# 3. 读取附件1，并保持与官方模板逐行一一对应
# ============================================================
def read_input_data(path: Path) -> dict[str, np.ndarray]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    header_map = {
        str(value).strip(): index + 1
        for index, value in enumerate(headers)
        if value is not None
    }

    required = ["时间", "电价", "小区负载", "光伏发电预测功率"]
    missing = [name for name in required if name not in header_map]
    if missing:
        wb.close()
        raise ValueError(f"数据文件缺少列：{missing}\n实际表头为：{headers}")

    raw_minutes: list[int] = []
    price: list[float] = []
    load_power: list[float] = []
    pv_power: list[float] = []

    for row in range(2, ws.max_row + 1):
        time_value = ws.cell(row, header_map["时间"]).value
        if time_value is None:
            continue

        values = [
            ws.cell(row, header_map["电价"]).value,
            ws.cell(row, header_map["小区负载"]).value,
            ws.cell(row, header_map["光伏发电预测功率"]).value,
        ]
        if any(value is None for value in values):
            wb.close()
            raise ValueError(f"第{row}行存在空数据")

        raw_minutes.append(excel_time_to_absolute_minutes(time_value))
        price.append(float(values[0]))
        load_power.append(float(values[1]))
        pv_power.append(float(values[2]))

    wb.close()

    if len(raw_minutes) != N:
        raise ValueError(f"附件1应有{N}个十分钟数据项，实际读取到{len(raw_minutes)}个")

    absolute_minutes = unwrap_time_sequence(raw_minutes)
    if len(np.unique(absolute_minutes)) != N:
        raise ValueError("时间序列中存在重复时刻")

    price_array = np.asarray(price, dtype=float)
    load_power_array = np.asarray(load_power, dtype=float)
    pv_power_array = np.asarray(pv_power, dtype=float)
    for name, array in {
        "电价": price_array,
        "小区负载": load_power_array,
        "光伏发电预测功率": pv_power_array,
    }.items():
        if not np.all(np.isfinite(array)):
            raise ValueError(f"{name}中存在非有限数")

    if np.any(price_array < 0) or np.any(load_power_array < 0) or np.any(pv_power_array < 0):
        raise ValueError("电价、负载和光伏功率均应为非负数")

    return {
        "absolute_minutes": absolute_minutes,
        "price": price_array,
        "load_power": load_power_array,
        "pv_power": pv_power_array,
        "load_energy": load_power_array * DELTA_T,
        "pv_energy": pv_power_array * DELTA_T,
    }


# ============================================================
# 4. 问题1确定性MILP
# ============================================================
def validate_solution(solution: dict[str, np.ndarray | float]) -> dict[str, float]:
    G = np.asarray(solution["G"])
    C = np.asarray(solution["C"])
    D = np.asarray(solution["D"])
    W = np.asarray(solution["W"])
    E = np.asarray(solution["E"])
    z = np.asarray(solution["z"])
    load_energy = np.asarray(solution["load_energy"])
    pv_energy = np.asarray(solution["pv_energy"])

    balance_residual = G + pv_energy + D - load_energy - C - W
    soc_residual = E[1:] - E[:-1] - ETA_C * C + D / ETA_D

    diagnostics = {
        "max_balance_residual": float(np.max(np.abs(balance_residual))),
        "max_soc_residual": float(np.max(np.abs(soc_residual))),
        "terminal_soc_residual": float(abs(E[-1] - E[0])),
        "max_charge_discharge_product": float(np.max(C * D)),
        "max_binary_residual": float(np.max(np.abs(z - np.rint(z)))),
        "min_soc": float(np.min(E)),
        "max_soc": float(np.max(E)),
    }

    if diagnostics["max_balance_residual"] > CHECK_TOL:
        raise AssertionError(f"电量平衡残差过大：{diagnostics}")
    if diagnostics["max_soc_residual"] > CHECK_TOL:
        raise AssertionError(f"SOC递推残差过大：{diagnostics}")
    if diagnostics["terminal_soc_residual"] > CHECK_TOL:
        raise AssertionError(f"日初日末SOC不一致：{diagnostics}")
    if diagnostics["max_charge_discharge_product"] > CHECK_TOL:
        raise AssertionError(f"存在同时充放电：{diagnostics}")
    if diagnostics["max_binary_residual"] > CHECK_TOL:
        raise AssertionError(f"二元变量不满足整数性：{diagnostics}")
    if diagnostics["min_soc"] < E_MIN - CHECK_TOL or diagnostics["max_soc"] > E_MAX + CHECK_TOL:
        raise AssertionError(f"SOC越界：{diagnostics}")
    if np.min(np.r_[G, C, D, W]) < -CHECK_TOL:
        raise AssertionError("存在显著负决策变量")
    if np.max(C) > Q_MAX + CHECK_TOL or np.max(D) > Q_MAX + CHECK_TOL:
        raise AssertionError("充放电量超过单时段上限")

    return diagnostics


def solve_one_scenario(
    initial_energy: float,
    price: np.ndarray,
    load_energy: np.ndarray,
    pv_energy: np.ndarray,
) -> dict[str, np.ndarray | float | dict[str, float]]:
    """求解给定日初储电量下的全天计划。C、D均为母线侧电量。"""
    if not (E_MIN <= initial_energy <= E_MAX):
        raise ValueError(f"初始储电量必须位于[{E_MIN}, {E_MAX}] kWh")

    # 变量顺序：G(N), C(N), D(N), W(N), E(N+1), z(N)
    G0 = 0
    C0 = N
    D0 = 2 * N
    W0 = 3 * N
    EIDX0 = 4 * N
    Z0 = EIDX0 + (N + 1)
    nvar = Z0 + N

    objective = np.zeros(nvar, dtype=float)
    objective[G0:G0 + N] = price

    # N条母线电量平衡 + N条SOC状态递推
    Aeq = lil_matrix((2 * N, nvar), dtype=float)
    beq = np.zeros(2 * N, dtype=float)

    for t in range(N):
        # G_t + V_t + D_t = L_t + C_t + W_t
        Aeq[t, G0 + t] = 1.0
        Aeq[t, C0 + t] = -1.0
        Aeq[t, D0 + t] = 1.0
        Aeq[t, W0 + t] = -1.0
        beq[t] = load_energy[t] - pv_energy[t]

        # E_{t+1} = E_t + eta_c*C_t - D_t/eta_d
        row = N + t
        Aeq[row, EIDX0 + t + 1] = 1.0
        Aeq[row, EIDX0 + t] = -1.0
        Aeq[row, C0 + t] = -ETA_C
        Aeq[row, D0 + t] = 1.0 / ETA_D

    constraints: list[LinearConstraint] = [
        LinearConstraint(csr_matrix(Aeq), beq, beq)
    ]

    # z_t=1允许充电；z_t=0允许放电
    Aub = lil_matrix((2 * N, nvar), dtype=float)
    upper = np.zeros(2 * N, dtype=float)
    for t in range(N):
        # C_t <= Q_MAX*z_t
        Aub[t, C0 + t] = 1.0
        Aub[t, Z0 + t] = -Q_MAX

        # D_t <= Q_MAX*(1-z_t)
        Aub[N + t, D0 + t] = 1.0
        Aub[N + t, Z0 + t] = Q_MAX
        upper[N + t] = Q_MAX

    constraints.append(
        LinearConstraint(
            csr_matrix(Aub),
            -np.inf * np.ones(2 * N),
            upper,
        )
    )

    lower_bounds = np.zeros(nvar, dtype=float)
    upper_bounds = np.full(nvar, np.inf, dtype=float)

    upper_bounds[C0:C0 + N] = Q_MAX
    upper_bounds[D0:D0 + N] = Q_MAX

    lower_bounds[EIDX0:EIDX0 + N + 1] = E_MIN
    upper_bounds[EIDX0:EIDX0 + N + 1] = E_MAX
    # 日初电量固定，且问题1要求日末与日初相同
    lower_bounds[EIDX0] = upper_bounds[EIDX0] = initial_energy
    lower_bounds[EIDX0 + N] = upper_bounds[EIDX0 + N] = initial_energy

    upper_bounds[Z0:Z0 + N] = 1.0
    integrality = np.zeros(nvar, dtype=int)
    integrality[Z0:Z0 + N] = 1

    result = milp(
        c=objective,
        integrality=integrality,
        bounds=Bounds(lower_bounds, upper_bounds),
        constraints=constraints,
        options={"time_limit": 120, "mip_rel_gap": 1e-9},
    )

    if not result.success:
        raise RuntimeError(
            f"E0={initial_energy} kWh 求解失败。\n求解器信息：{result.message}"
        )

    x = result.x
    G = x[G0:G0 + N].copy()
    C = x[C0:C0 + N].copy()
    D = x[D0:D0 + N].copy()
    W = x[W0:W0 + N].copy()
    E = x[EIDX0:EIDX0 + N + 1].copy()
    z = x[Z0:Z0 + N].copy()

    for array in (G, C, D, W, E):
        array[np.abs(array) < ZERO_TOL] = 0.0
    z = np.rint(z).astype(int)

    solution: dict[str, np.ndarray | float | dict[str, float]] = {
        "E0": float(initial_energy),
        "G": G,
        "C": C,
        "D": D,
        "W": W,
        "E": E,
        "z": z,
        "price": price,
        "load_energy": load_energy,
        "pv_energy": pv_energy,
        "total_grid": float(np.sum(G)),
        "total_cost": float(np.dot(price, G)),
        "total_charge": float(np.sum(C)),
        "total_discharge": float(np.sum(D)),
        "total_curtailment": float(np.sum(W)),
    }
    solution["diagnostics"] = validate_solution(solution)
    return solution


# ============================================================
# 5. 输出官方 result1 模板
# ============================================================
BLOCK_NAMES = [
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
]

BLOCKS = [
    (0, 24),
    (24, 48),
    (48, 72),
    (72, 96),
    (96, 120),
    (120, 144),
]

TABLE1_TIMES = [
    "10:00-10:10",
    "12:00-12:10",
    "14:00-14:10",
    "16:00-16:10",
    "18:00-18:10",
    "20:00-20:10",
]


def build_time_index(absolute_minutes: np.ndarray) -> dict[int, int]:
    return {int(minute): index for index, minute in enumerate(absolute_minutes)}


def fill_result1(
    solution: dict[str, np.ndarray | float | dict[str, float]],
    template_file: Path,
    output_path: Path,
    time_to_index: dict[int, int],
) -> None:
    """按官方模板时间标签匹配购电量，并填入母线侧充、放电量。"""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_file, output_path)
    wb = load_workbook(output_path)

    required_sheets = {"计划购电量", "充放电量"}
    if not required_sheets.issubset(wb.sheetnames):
        wb.close()
        raise ValueError("result1模板必须包含“计划购电量”和“充放电量”工作表")

    ws_buy = wb["计划购电量"]
    ws_battery = wb["充放电量"]
    G = np.asarray(solution["G"])
    C = np.asarray(solution["C"])
    D = np.asarray(solution["D"])
    E = np.asarray(solution["E"])

    matched_rows = 0
    for row in range(2, ws_buy.max_row + 1):
        label = ws_buy.cell(row, 1).value
        if label is None:
            continue
        start_minute = parse_interval_start(label)
        if start_minute not in time_to_index:
            wb.close()
            raise ValueError(
                f"模板时间段{label}的起点{absolute_minute_text(start_minute)}"
                "在附件1中没有对应数据"
            )
        index = time_to_index[start_minute]
        ws_buy.cell(row, 2).value = round(float(G[index]), 6)
        ws_buy.cell(row, 2).number_format = "0.000000"
        matched_rows += 1

    if matched_rows != N:
        wb.close()
        raise ValueError(f"计划购电量工作表应匹配{N}行，实际匹配{matched_rows}行")

    for row, (start, end) in enumerate(BLOCKS, start=2):
        ws_battery.cell(row, 2).value = round(float(np.sum(C[start:end])), 6)
        ws_battery.cell(row, 3).value = round(float(np.sum(D[start:end])), 6)
        ws_battery.cell(row, 2).number_format = "0.000000"
        ws_battery.cell(row, 3).number_format = "0.000000"

    ws_battery["E2"] = round(float(E[0]), 6)
    ws_battery["E3"] = round(float(E[-1]), 6)
    ws_battery["E2"].number_format = "0.000000"
    ws_battery["E3"].number_format = "0.000000"

    # 模板数值列采用默认宽度，六位小数会显示不全；仅扩宽受影响列。
    for worksheet, columns in [
        (ws_buy, ("B",)),
        (ws_battery, ("B", "C", "E")),
    ]:
        for column in columns:
            current_width = worksheet.column_dimensions[column].width or 8.43
            worksheet.column_dimensions[column].width = max(current_width, 15.0)

    wb.save(output_path)
    wb.close()


def create_summary_excel(
    solutions: dict[int, dict[str, np.ndarray | float | dict[str, float]]],
    output_path: Path,
) -> None:
    """生成黑白样式的初始SOC敏感性汇总表。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "初始SOC敏感性"

    headers = [
        "初始储电量E0(kWh)",
        "24:00储电量(kWh)",
        "全天购电量(kWh)",
        "全天购电费(元)",
        "母线侧充电量(kWh)",
        "母线侧放电量(kWh)",
        "未利用富余电量(kWh)",
    ]
    ws.append(headers)

    for initial_energy in E0_LIST:
        solution = solutions[initial_energy]
        E = np.asarray(solution["E"])
        ws.append([
            initial_energy,
            float(E[-1]),
            float(solution["total_grid"]),
            float(solution["total_cost"]),
            float(solution["total_charge"]),
            float(solution["total_discharge"]),
            float(solution["total_curtailment"]),
        ])

    thin_black = Side(style="thin", color="000000")
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="000000")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=thin_black, bottom=thin_black)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="000000")
            cell.number_format = "0.000000"
            cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [22, 22, 22, 22, 24, 24, 26]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()


def create_text_result(
    solutions: dict[int, dict[str, np.ndarray | float | dict[str, float]]],
    output_path: Path,
    time_to_index: dict[int, int],
) -> None:
    lines = [
        "问题1：计划购电策略计算结果",
        "=" * 72,
        "离散步长：10分钟；共144个数据项，并与官方模板144行按顺序/绝对时刻对应。",
        "时间解析保留跨日偏移：0:00+1=1440分钟。",
        "目标函数：最小化全天计划购电费用 Σ p_t G_t。",
        "",
        "储能变量口径：",
        "  C_t：母线侧送入储能的充电量，SOC增加0.9*C_t。",
        "  D_t：储能向母线提供的有效放电量，SOC减少D_t/0.9。",
        "  C_t、D_t每10分钟均不超过5000/6=833.333333 kWh。",
        "  z_t：充放电互斥二元变量；W_t：未利用富余电量。",
        "",
    ]

    for initial_energy in E0_LIST:
        solution = solutions[initial_energy]
        E = np.asarray(solution["E"])
        C = np.asarray(solution["C"])
        D = np.asarray(solution["D"])
        G = np.asarray(solution["G"])
        diagnostics = solution["diagnostics"]

        lines.extend([
            f"【初始储电量 E0 = {initial_energy} kWh】",
            f"全天购电量：{float(solution['total_grid']):.6f} kWh",
            f"全天购电费：{float(solution['total_cost']):.6f} 元",
            f"全天母线侧充电量：{float(solution['total_charge']):.6f} kWh",
            f"全天母线侧放电量：{float(solution['total_discharge']):.6f} kWh",
            f"全天未利用富余电量：{float(solution['total_curtailment']):.6f} kWh",
            "",
            "表1指定时段购电量：",
        ])

        for label in TABLE1_TIMES:
            minute = parse_interval_start(label)
            index = time_to_index[minute]
            lines.append(f"  {label}：{G[index]:.6f} kWh")

        lines.append("")
        lines.append("表2指定时段储能充放电量（均为母线侧）：")
        for name, (start, end) in zip(BLOCK_NAMES, BLOCKS):
            lines.append(
                f"  {name}：充电 {np.sum(C[start:end]):.6f} kWh，"
                f"放电 {np.sum(D[start:end]):.6f} kWh"
            )

        lines.extend([
            f"0:00储电量：{E[0]:.6f} kWh",
            f"24:00储电量：{E[-1]:.6f} kWh",
            "自动校验："
            f"平衡残差={diagnostics['max_balance_residual']:.3e}，"
            f"SOC残差={diagnostics['max_soc_residual']:.3e}，"
            f"终端残差={diagnostics['terminal_soc_residual']:.3e}",
            "-" * 72,
        ])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(lines), encoding="utf-8-sig")


# ============================================================
# 6. 主程序
# ============================================================
def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="问题1：微网确定性储能调度MILP")
    parser.add_argument("--data", type=Path, default=None, help="附件1.xlsx路径")
    parser.add_argument("--template", type=Path, default=None, help="result1.xlsx模板路径")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="输出文件夹，默认：问题1输出",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    data_file = find_existing_file(args.data, DATA_CANDIDATES, "数据文件")
    template_file = find_existing_file(args.template, TEMPLATE_CANDIDATES, "result1模板")
    output_dir: Path = args.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)

    data = read_input_data(data_file)
    time_to_index = build_time_index(data["absolute_minutes"])

    print("=" * 72)
    print("问题1 微网全天计划购电策略（统一母线侧口径）")
    print("=" * 72)
    print(f"数据文件：{data_file}")
    print(f"模板文件：{template_file}")
    print(f"输出目录：{output_dir}")
    print(
        "时间范围："
        f"{absolute_minute_text(int(data['absolute_minutes'][0]))} 至 "
        f"{absolute_minute_text(int(data['absolute_minutes'][-1]))}"
    )
    print()

    solutions: dict[int, dict[str, np.ndarray | float | dict[str, float]]] = {}
    for initial_energy in E0_LIST:
        print(f"正在求解 E0={initial_energy} kWh ...")
        solution = solve_one_scenario(
            initial_energy=initial_energy,
            price=data["price"],
            load_energy=data["load_energy"],
            pv_energy=data["pv_energy"],
        )
        solutions[initial_energy] = solution
        print(
            f"  完成：全天购电量={float(solution['total_grid']):.6f} kWh，"
            f"购电费={float(solution['total_cost']):.6f} 元"
        )

    for initial_energy in E0_LIST:
        fill_result1(
            solution=solutions[initial_energy],
            template_file=template_file,
            output_path=output_dir / f"result1_E0_{initial_energy}.xlsx",
            time_to_index=time_to_index,
        )

    fill_result1(
        solution=solutions[BASE_E0],
        template_file=template_file,
        output_path=output_dir / "result1.xlsx",
        time_to_index=time_to_index,
    )

    create_summary_excel(
        solutions,
        output_dir / "问题1_初始储电量敏感性汇总.xlsx",
    )
    create_text_result(
        solutions,
        output_dir / "问题1_文字结果.txt",
        time_to_index,
    )

    print()
    print("=" * 72)
    print(f"E0={BASE_E0} kWh 基准结果")
    print("=" * 72)
    base_solution = solutions[BASE_E0]
    print(f"全天购电量：{float(base_solution['total_grid']):.6f} kWh")
    print(f"全天购电费：{float(base_solution['total_cost']):.6f} 元")
    print(f"母线侧充电量：{float(base_solution['total_charge']):.6f} kWh")
    print(f"母线侧放电量：{float(base_solution['total_discharge']):.6f} kWh")
    print(f"结果文件：{output_dir / 'result1.xlsx'}")


if __name__ == "__main__":
    main()
