from __future__ import annotations

import csv
import io
import json
import os
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse

try:
    from .model_service import solve_scenario, solve_forecast, solve_q3, solve_q4
except ImportError:  # uvicorn server:app --app-dir backend
    from model_service import solve_scenario, solve_forecast, solve_q3, solve_q4

app = FastAPI(title="GridScope Scenario API", version="1.0.0")
_default_origins = [
    "http://127.0.0.1:4173",
    "http://localhost:4173",
    "https://yigedashuaige-hi.github.io",
]
_extra_origins = [origin.strip() for origin in os.getenv("GRIDSCOPE_ALLOWED_ORIGINS", "").split(",") if origin.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=list(dict.fromkeys(_default_origins + _extra_origins)),
    allow_credentials=False,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["Content-Type"],
)


@app.get("/api/health")
def health() -> dict[str, Any]:
    return {
        "ok": True,
        "status": "online",
        "service": "GridScope Scenario API",
        "solver": "q1.py / q2.py / q3.py / q4.py · scipy HiGHS",
        "version": app.version,
    }


@app.post("/api/scenario/solve")
def scenario(payload: dict[str, Any]) -> dict[str, Any]:
    try:
        return solve_scenario(payload)
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.post("/api/forecast/solve")
def forecast(payload: dict[str, Any]) -> dict[str, Any]:
    try:
        return solve_forecast(payload)
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


def _parse_history_rows(raw: bytes, suffix: str) -> list[dict[str, Any]]:
    required = {"date", "time", "load_kw", "pv_kw", "price_yuan_kwh"}
    if suffix == ".csv":
        try:
            rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
        except UnicodeDecodeError as exc:
            raise ValueError("history CSV must be UTF-8 encoded") from exc
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        values = list(wb.active.iter_rows(values_only=True)); wb.close()
        if not values:
            raise ValueError("history workbook is empty")
        headers = [str(value).strip() for value in values[0]]
        rows = [dict(zip(headers, row)) for row in values[1:]]
    else:
        raise ValueError("upload a .csv or .xlsx file")
    if not rows or not required.issubset(rows[0].keys()):
        raise ValueError("forecast history requires date,time,load_kw,pv_kw,price_yuan_kwh columns")
    normalized = []
    for row in rows:
        item = {"date": str(row["date"])[:10], "time": str(row["time"])}
        for key in ("load_kw", "pv_kw", "price_yuan_kwh"):
            item[key] = float(row[key])
        normalized.append(item)
    return normalized


@app.post("/api/forecast/solve-file")
async def forecast_file(file: UploadFile = File(...), date: str = "", initial_soc_percent: float = 50.0, alpha: float = 0.70) -> dict[str, Any]:
    suffix = Path(file.filename or "").suffix.lower()
    try:
        rows = _parse_history_rows(await file.read(), suffix)
        target = date or sorted({row["date"] for row in rows})[-1]
        return solve_forecast({"date": target, "initial_soc_percent": initial_soc_percent, "alpha": alpha, "history_rows": rows, "uploaded": True, "history_days": len({row["date"] for row in rows})})
    except (ValueError, RuntimeError, TypeError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.post("/api/q3/solve")
def q3_online(payload: dict[str, Any]) -> dict[str, Any]:
    try:
        return solve_q3(payload)
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.post("/api/q4/solve")
def q4_online(payload: dict[str, Any]) -> dict[str, Any]:
    try:
        return solve_q4(payload)
    except (ValueError, RuntimeError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


def _parse_expert_csv(raw: bytes) -> dict[str, list[float]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV must be UTF-8 encoded") from exc
    rows = list(csv.DictReader(io.StringIO(text)))
    required = {"time", "load_kw", "pv_kw", "price_yuan_kwh"}
    if not rows or not required.issubset(rows[0].keys()):
        raise ValueError("file must contain time, load_kw, pv_kw, price_yuan_kwh columns")
    if len(rows) != 144:
        raise ValueError("expert file must contain exactly 144 rows")
    return {key: [float(row[key]) for row in rows] for key in ("load_kw", "pv_kw", "price_yuan_kwh")}


@app.post("/api/scenario/solve-file")
async def scenario_file(file: UploadFile = File(...), initial_soc_percent: float = 50.0, alpha: float = 0.70) -> dict[str, Any]:
    suffix = Path(file.filename or "").suffix.lower()
    raw = await file.read()
    try:
        if suffix == ".csv":
            values = _parse_expert_csv(raw)
        elif suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            rows = list(wb.active.iter_rows(values_only=True))
            headers = [str(x).strip() for x in rows[0]]
            indices = {name: headers.index(name) for name in ("time", "load_kw", "pv_kw", "price_yuan_kwh")}
            if len(rows) != 145:
                raise ValueError("expert workbook must contain a header plus 144 rows")
            values = {name: [float(row[indices[name]]) for row in rows[1:]] for name in indices if name != "time"}
            wb.close()
        else:
            raise ValueError("upload a .csv or .xlsx file")
        return solve_scenario({"mode": "expert", **values, "initial_soc_percent": initial_soc_percent, "alpha": alpha})
    except (ValueError, KeyError, IndexError, TypeError, RuntimeError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.exception_handler(Exception)
async def unhandled(_request: Any, exc: Exception) -> JSONResponse:
    return JSONResponse(status_code=500, content={"detail": f"solver error: {exc}"})
