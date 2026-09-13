#!/usr/bin/env python3
"""Export the read-only web snapshot from the project's final inputs/results.

This exporter never writes into the official or result directories.  It keeps
the full 144-slot source curves and preserves emergency purchases as the
contiguous events recorded by the result workbooks (the workbooks do not carry
their hidden per-slot split).
"""
from __future__ import annotations

import argparse
import csv
import importlib.util
import sys
import json
import math
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SLOTS = 144
DELTA_H = 1 / 6
ETA_C = 0.9
ETA_D = 0.9


def number(value: Any, field: str) -> float:
    try:
        out = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field} contains a non-numeric value: {value!r}") from exc
    if not math.isfinite(out):
        raise ValueError(f"{field} contains a non-finite value")
    return out


def read_actual(path: Path) -> tuple[list[str], dict[str, dict[str, list[float]]]]:
    records: dict[str, dict[str, list[float]]] = {}
    slot_labels: list[str] = []
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = csv.DictReader(stream)
        required = {"日期", "时间", "小区负载（kW）", "光伏实际功率（kW）", "实际电价（元/kWh）"}
        missing = required - set(rows.fieldnames or [])
        if missing:
            raise ValueError(f"actual CSV missing columns: {sorted(missing)}")
        for row in rows:
            date = str(row["日期"])
            target = records.setdefault(date, {"load_kw": [], "pv_kw": [], "price_yuan_kwh": [], "time": []})
            target["time"].append(str(row["时间"]))
            target["load_kw"].append(number(row["小区负载（kW）"], "load"))
            target["pv_kw"].append(number(row["光伏实际功率（kW）"], "pv"))
            target["price_yuan_kwh"].append(number(row["实际电价（元/kWh）"], "price"))
        for date, row in records.items():
            if len(row["load_kw"]) != SLOTS:
                raise ValueError(f"{date} has {len(row['load_kw'])} slots, expected {SLOTS}")
            if not slot_labels:
                slot_labels = row["time"]
    return slot_labels, records


def read_causal_forecast(site_root: Path, dates: list[str]) -> dict[str, dict[str, Any]]:
    """Export the final q2 predictor path without solving any control problem."""
    q2_path = site_root / "model" / "q2.py"
    data_path = site_root / "source" / "processed" / "附件2_处理后.xlsx"
    price_path = site_root / "source" / "processed" / "附件1_处理后.xlsx"
    if not (q2_path.exists() and data_path.exists() and price_path.exists()):
        return {}
    spec = importlib.util.spec_from_file_location("gridscope_export_q2", q2_path)
    if spec is None or spec.loader is None:
        return {}
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    model_data = module.load_inputs(data_path, price_path)
    forecaster = module.CausalForecaster(model_data)
    methods, records, selected, _ = module.build_daily_forecast_policy(model_data, forecaster)
    available = set(dates)
    output: dict[str, dict[str, Any]] = {}
    for index, date_value in enumerate(model_data.dates.strftime("%Y-%m-%d")):
        if date_value not in available:
            continue
        output[date_value] = {
            "selected_predictor": methods[index],
            "selection_rule": records[index]["selection_rule"],
            "history_days": records[index]["history_days"],
            "center_net_kw": (selected[index] / module.DELTA_H).round(6).tolist(),
        }
    return output


def read_causal_price_forecast(site_root: Path, dates: list[str], causal: dict[str, dict[str, Any]]) -> dict[str, dict[str, Any]]:
    q4_path = site_root / "model" / "q4.py"
    q2_path = site_root / "model" / "q2.py"
    if not (q4_path.exists() and q2_path.exists()):
        return {}
    q2_spec = importlib.util.spec_from_file_location("gridscope_price_q2", q2_path)
    q4_spec = importlib.util.spec_from_file_location("gridscope_export_q4", q4_path)
    if q2_spec is None or q2_spec.loader is None or q4_spec is None or q4_spec.loader is None:
        return {}
    q2 = importlib.util.module_from_spec(q2_spec); sys.modules[q2_spec.name] = q2; q2_spec.loader.exec_module(q2)
    q4 = importlib.util.module_from_spec(q4_spec); sys.modules[q4_spec.name] = q4; q4_spec.loader.exec_module(q4)
    processed = site_root / "source" / "processed"
    data = q2.load_inputs(processed / "附件2_处理后.xlsx", processed / "附件1_处理后.xlsx")
    actual = q4.load_actual_prices(processed / "附件4_处理后.xlsx", data.dates)
    net_forecaster = q2.CausalForecaster(data)
    _, _, selected_predictions, _ = q2.build_daily_forecast_policy(data, net_forecaster)
    forecaster = q4.CausalPriceForecaster(actual, data.dates, data.net - selected_predictions)
    output: dict[str, dict[str, Any]] = {}
    for day, date_value in enumerate(data.dates.strftime("%Y-%m-%d")):
        if date_value not in dates:
            continue
        output[date_value] = {
            "model": forecaster.model_by_day[day],
            "forecast_yuan_kwh": forecaster.forecast_before_slot(day, 0, SLOTS).round(8).tolist(),
            "baseline_yuan_kwh": forecaster.baseline[day].round(8).tolist(),
            "actual_yuan_kwh": actual[day].round(8).tolist(),
        }
    return output


def _date_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    return text[:10] if text else None


def _vector(row: tuple[Any, ...], start: int = 1) -> list[float]:
    values = [number(x or 0, "result value") for x in row[start:start + SLOTS]]
    if len(values) != SLOTS:
        raise ValueError(f"result row has {len(values)} slots, expected {SLOTS}")
    return values


def read_battery(ws: Any) -> dict[str, dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))[1:]
    output: dict[str, dict[str, Any]] = {}
    current_date: str | None = None
    bucket: list[tuple[Any, ...]] = []
    def flush() -> None:
        nonlocal bucket, current_date
        if not current_date or len(bucket) < 6:
            bucket = []
            return
        charge = [number(row[2] or 0, "charge") for row in bucket[:6]]
        discharge = [number(row[3] or 0, "discharge") for row in bucket[:6]]
        start = next((number(row[5], "soc") for row in bucket if row[5] is not None), 6000.0)
        soc = [start]
        for c, d in zip(charge, discharge):
            soc.append(soc[-1] + ETA_C * c - d / ETA_D)
        output[current_date] = {
            "charge_kwh_4h": charge,
            "discharge_kwh_4h": discharge,
            "soc_kwh": [round(x, 6) for x in soc],
        }
        bucket = []
    for row in rows:
        date = _date_text(row[0])
        if date:
            if current_date and date != current_date:
                flush()
            current_date = date
        bucket.append(row)
    flush()
    return output


def read_strategy(path: Path) -> dict[str, dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    plan_rows = list(wb["计划购电量"].iter_rows(values_only=True))[1:]
    adjusted_rows = list(wb["调整购电量"].iter_rows(values_only=True))[1:] if "调整购电量" in wb.sheetnames else []
    battery = read_battery(wb["充放电量"])
    emergency_rows = list(wb["紧急购电量"].iter_rows(values_only=True))[1:]
    events: dict[str, list[dict[str, Any]]] = defaultdict(list)
    current_date: str | None = None
    for row in emergency_rows:
        date = _date_text(row[0])
        if date:
            current_date = date
        if current_date and row[1] not in (None, "无") and row[2] not in (None, ""):
            events[current_date].append({"period": str(row[1]), "energy_kwh": number(row[2], "emergency")})
    adjusted_by_date: dict[str, list[float]] = {}
    for row in adjusted_rows:
        date = _date_text(row[0])
        if date:
            adjusted_by_date[date] = _vector(row)
    result: dict[str, dict[str, Any]] = {}
    for row in plan_rows:
        date = _date_text(row[0])
        if not date:
            continue
        result[date] = {
            "plan_kwh": _vector(row),
            "adjusted_plan_kwh": adjusted_by_date.get(date),
            "emergency_events": events.get(date, []),
            "battery": battery.get(date),
        }
    wb.close()
    return result


def export(args: argparse.Namespace) -> dict[str, Any]:
    slots, actual = read_actual(args.actual)
    q2 = read_strategy(args.q2)
    q3 = read_strategy(args.q3)
    q42 = read_strategy(args.q42)
    q43 = read_strategy(args.q43)
    dates = sorted(date for date in actual if date >= "2025-02-01")
    site_root = Path(__file__).resolve().parents[1]
    causal = read_causal_forecast(site_root, dates)
    causal_price = read_causal_price_forecast(site_root, dates, causal)
    days: dict[str, Any] = {}
    for date in dates:
        days[date] = {
            "actual": actual[date],
            "causal": causal.get(date),
            "causal_price": causal_price.get(date),
            "strategies": {
                "Q2": q2.get(date), "Q3-M3": q3.get(date),
                "Q4-2": q42.get(date), "Q4-3": q43.get(date),
            },
        }
    payload = {
        "schema_version": "1.0",
        "generated_from": {
            "actual": args.actual.name, "q2": args.q2.name, "q3": args.q3.name,
            "q42": args.q42.name, "q43": args.q43.name,
        },
        "meta": {
            "slot_minutes": 10, "slots_per_day": SLOTS, "delta_h": DELTA_H,
            "eta_charge": ETA_C, "eta_discharge": ETA_D,
            "soc_min_kwh": 1200, "soc_max_kwh": 10800, "capacity_kwh": 12000,
            "date_start": dates[0], "date_end": dates[-1],
            "emergency_representation": "contiguous events from result workbook; no per-slot split is inferred",
            "official_results_read_only": True,
        },
        "slot_labels": slots,
        "days": days,
        "frozen_metrics": build_frozen_metrics(args, actual, q2, q3, q42, q43),
        "metric_sources": {
            "q1": "source/results/result1_问题一.xlsx / plan and battery sheets",
            "q2": "source/results/result2_问题二.xlsx / 144-slot plan and emergency event sheets",
            "q3": "source/results/result3_问题三.xlsx / M3 plan and adjustment sheets",
            "q4": "source/results/result4-2_问题四.xlsx and result4-3_问题四.xlsx",
            "event_costs": "Emergency events remain aggregated by workbook period; no per-slot split is inferred.",
        },
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    return payload


def _strategy_energy(strategy: dict[str, dict[str, Any]]) -> float:
    return float(sum(sum(row.get("plan_kwh") or []) for row in strategy.values() if row))


def _strategy_adjusted_energy(strategy: dict[str, dict[str, Any]]) -> float:
    return float(sum(sum(row.get("adjusted_plan_kwh") or []) for row in strategy.values() if row))


def _strategy_emergency(strategy: dict[str, dict[str, Any]]) -> float:
    return float(sum(event.get("energy_kwh", 0.0) for row in strategy.values() if row for event in row.get("emergency_events", [])))


def _q1_metrics(site_root: Path) -> dict[str, float]:
    path = site_root / "source" / "results" / "result1_问题一.xlsx"
    if not path.exists():
        return {}
    wb = load_workbook(path, read_only=True, data_only=True)
    plan = [number(row[1], "q1 plan") for row in list(wb["计划购电量"].iter_rows(values_only=True))[1:] if row[1] is not None]
    battery = list(wb["充放电量"].iter_rows(values_only=True))[1:]
    wb.close()
    price_path = site_root / "source" / "processed" / "附件1_处理后.xlsx"
    price_wb = load_workbook(price_path, read_only=True, data_only=True)
    prices = [number(row[1], "q1 price") for row in list(price_wb.active.iter_rows(values_only=True))[1:1 + SLOTS]]
    price_wb.close()
    return {"q1_plan_kwh": float(sum(plan)), "q1_cost_yuan": float(sum(p * g for p, g in zip(prices, plan))), "q1_charge_kwh": float(sum((row[1] or 0) for row in battery)), "q1_discharge_kwh": float(sum((row[2] or 0) for row in battery))}


def build_frozen_metrics(args: argparse.Namespace, actual: dict[str, dict[str, list[float]]], q2: dict, q3: dict, q42: dict, q43: dict) -> dict[str, float | None]:
    site_root = Path(__file__).resolve().parents[1]
    q1 = _q1_metrics(site_root)
    selected_alpha = 0.70
    q3_module_path = site_root / "model" / "q3.py"
    if q3_module_path.exists():
        spec = importlib.util.spec_from_file_location("gridscope_metric_q3", q3_module_path)
        if spec is not None and spec.loader is not None:
            module = importlib.util.module_from_spec(spec); sys.modules[spec.name] = module; spec.loader.exec_module(module)
            selected_alpha = float(getattr(module, "SELECTED_ALPHA", selected_alpha))
    q2_plan = _strategy_energy(q2); q3_plan = _strategy_adjusted_energy(q3); q42_plan = _strategy_energy(q42); q43_plan = _strategy_adjusted_energy(q43)
    def plan_cost(strategy: dict[str, dict[str, Any]], adjusted: bool = False) -> float:
        total = 0.0
        field = "adjusted_plan_kwh" if adjusted else "plan_kwh"
        for date, row in strategy.items():
            if not row or date not in actual:
                continue
            values = row.get(field) or []
            total += sum(float(a) * float(b) for a, b in zip(values, actual[date]["price_yuan_kwh"]))
        return float(total)
    return {
        **q1,
        "q2_plan_kwh": q2_plan, "q2_plan_cost_yuan": plan_cost(q2), "q2_emergency_kwh": _strategy_emergency(q2),
        "m3_plan_kwh": q3_plan, "m3_plan_cost_yuan": plan_cost(q3, True), "m3_emergency_kwh": _strategy_emergency(q3),
        "q4_2_plan_kwh": q42_plan, "q4_2_plan_cost_yuan": plan_cost(q42), "q4_2_emergency_kwh": _strategy_emergency(q42),
        "q4_3_plan_kwh": q43_plan, "q4_3_plan_cost_yuan": plan_cost(q43, True), "q4_3_emergency_kwh": _strategy_emergency(q43),
        "alpha": selected_alpha,
    }


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    processed = root / "source" / "processed"
    results = root / "source" / "results"
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--actual", type=Path, default=processed / "全年实际数据.csv")
    parser.add_argument("--q2", type=Path, default=results / "result2_问题二.xlsx")
    parser.add_argument("--q3", type=Path, default=results / "result3_问题三.xlsx")
    parser.add_argument("--q42", type=Path, default=results / "result4-2_问题四.xlsx")
    parser.add_argument("--q43", type=Path, default=results / "result4-3_问题四.xlsx")
    parser.add_argument("--output", type=Path, default=root / "data/web_data.json")
    args = parser.parse_args()
    payload = export(args)
    print(json.dumps({"output": str(args.output), "dates": len(payload["days"]), "slots": payload["meta"]["slots_per_day"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
